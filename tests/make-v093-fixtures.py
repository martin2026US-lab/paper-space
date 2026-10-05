from pathlib import Path
from docx import Document
from reportlab.pdfgen import canvas
base=Path(__file__).parent
doc=Document();doc.add_heading('Table rendering fixture',1);doc.add_paragraph('Synthetic local test data. No private manuscript content.')
t=doc.add_table(rows=5, cols=5);t.style='Table Grid'
for c,s in zip(t.rows[0].cells,['Cohort','Method','Score [95% CI]','Outcome','Area fraction']):c.text=s
for i,row in enumerate(t.rows[1:],1):
    for c,s in zip(row.cells,['Group A' if i<3 else 'Group B','Method '+str(i),'0.800 [0.750, 0.850]','0.100 [0.050, 0.200]','0.025 [0.010, 0.050]']):c.text=s
t.cell(1,0).merge(t.cell(2,0));t.cell(3,3).merge(t.cell(3,4));doc.add_paragraph('Table footnote: all values above are synthetic.')
doc.save(base/'table-fixture.docx')
c=canvas.Canvas(str(base/'outline-fixture.pdf'),pagesize=(612,792))
for i in range(1,17):
    c.setFont('Helvetica',20);c.drawString(70,700,'Synthetic PDF page '+str(i))
    if i in (1,9,11,14,15,16):
        title={1:'Introduction',9:'References',11:'Tables',14:'Figure 1',15:'Figure 2',16:'Figure 3'}[i];c.bookmarkPage('p'+str(i));c.addOutlineEntry(title,'p'+str(i),0)
    c.showPage()
c.save()
print('Created synthetic DOCX table and 16-page outline fixtures')
